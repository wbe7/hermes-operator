"""Exercise installed tools as the unprivileged agent, with no external network."""
import functools
import http.server
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import threading
import unittest


def command(*args, timeout=120):
    result = subprocess.run(args, text=True, stdout=subprocess.PIPE,
                            stderr=subprocess.STDOUT, timeout=timeout)
    if result.returncode:
        raise AssertionError(f"{args[0]} exited {result.returncode}: {result.stdout[-6000:]}")
    return result.stdout


class QuietHandler(http.server.SimpleHTTPRequestHandler):
    def log_message(self, *args):
        pass


class AgentTools(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory(prefix="hermes-tools-")
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)

    def pdf_text(self, path):
        return command("pdftotext", str(path), "-")

    def office_pdf(self, path):
        command("libreoffice", "-env:UserInstallation=" + (self.root / "lo-profile").as_uri(),
                "--headless", "--convert-to", "pdf", "--outdir", str(self.root), str(path))
        result = self.root / (path.stem + ".pdf")
        self.assertTrue(result.is_file(), "LibreOffice did not create a PDF")
        return result

    def test_browser_python_and_hermes_cli(self):
        from playwright.sync_api import sync_playwright
        from pypdf import PdfReader

        (self.root / "index.html").write_text(
            '<!doctype html><title>Hermes browser smoke</title>'
            '<meta charset="utf-8"><h1>Browser ready</h1><p>Привет, мир</p>')
        server = http.server.ThreadingHTTPServer(("127.0.0.1", 0),
            functools.partial(QuietHandler, directory=str(self.root)))
        self.addCleanup(server.server_close)
        self.addCleanup(server.shutdown)
        threading.Thread(target=server.serve_forever, daemon=True).start()
        url = f"http://127.0.0.1:{server.server_port}/index.html"
        with sync_playwright() as p:
            # Match the upstream container browser flags. No host IPC/capability needed.
            browser = p.chromium.launch(args=["--no-sandbox", "--disable-dev-shm-usage"])
            page = browser.new_page()
            page.goto(url)
            self.assertEqual(page.title(), "Hermes browser smoke")
            page.screenshot(path=str(self.root / "browser.png"))
            page.pdf(path=str(self.root / "browser.pdf"))
            browser.close()
        self.assertIn("Browser ready", PdfReader(self.root / "browser.pdf").pages[0].extract_text())
        session = "hermes-tools-" + str(os.getpid())
        prefix = ("agent-browser", "--session", session)
        try:
            command(*prefix, "open", url)
            self.assertIn("Hermes browser smoke", command(*prefix, "eval", "document.title"))
        finally:
            command(*prefix, "close")

    def test_hermes_browser_exec(self):
        sys.path.insert(0, "/opt/hermes")
        from tools.browser_use_cli import browser_exec

        code = ('new_tab("about:blank")\n'
                'js("document.body.innerHTML = \'<h1>Hermes CDP ready</h1>\'")\n'
                'print(js("document.body.innerText"))\n'
                'print(capture_screenshot())\n')
        result = browser_exec(code, session="hermes-tools-smoke", timeout_s=90)
        text = result if isinstance(result, str) else json.dumps(result)
        self.assertIn("Hermes CDP ready", text, text[:1500])
        self.assertNotIn('"success": false', text, text[:1500])

    def test_docx_templates_and_pandoc(self):
        from docx import Document
        from docxtpl import DocxTemplate

        path = self.root / "document.docx"
        document = Document()
        document.add_heading("Hermes document", 0)
        document.add_paragraph("Hello {{ customer }}")
        document.save(path)
        template = DocxTemplate(path)
        template.render({"customer": "Maria"})
        template.save(path)
        self.assertIn("Hello Maria", self.pdf_text(self.office_pdf(path)))
        source = self.root / "source.md"
        source.write_text("# Pandoc document\n\nПривет, мир\n")
        target = self.root / "pandoc.docx"
        command("pandoc", str(source), "-o", str(target))
        self.assertIn("Привет, мир", "\n".join(p.text for p in Document(target).paragraphs))

    def test_spreadsheets_and_formula_recalculation(self):
        import openpyxl
        import xlsxwriter

        path = self.root / "sheet.xlsx"
        with xlsxwriter.Workbook(path) as book:
            sheet = book.add_worksheet("Report")
            sheet.write("A1", 12)
            sheet.write("A2", 30)
            sheet.write_formula("A3", "=SUM(A1:A2)")
        converted = self.root / "recalculated"
        converted.mkdir()
        command("hermes-recalculate", str(path), str(converted / path.name))
        book = openpyxl.load_workbook(converted / path.name, data_only=True)
        self.assertEqual(book["Report"]["A3"].value, 42)
        book.close()
        self.assertIn("42", self.pdf_text(self.office_pdf(converted / path.name)))

    def test_presentation_to_pdf(self):
        from pptx import Presentation

        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[0])
        slide.shapes.title.text = "Hermes presentation"
        path = self.root / "slides.pptx"
        presentation.save(path)
        self.assertIn("Hermes presentation", self.pdf_text(self.office_pdf(path)))

    def test_pdf_creation_extraction_merge_and_raster(self):
        import pdfplumber
        from pdf2image import convert_from_path
        from pypdf import PdfReader, PdfWriter
        from reportlab.pdfgen import canvas

        source = self.root / "source.pdf"
        pdf = canvas.Canvas(str(source))
        pdf.drawString(72, 720, "PDF tools ready")
        pdf.save()
        with pdfplumber.open(source) as document:
            self.assertIn("PDF tools ready", document.pages[0].extract_text())
        writer = PdfWriter()
        writer.append(source)
        writer.append(source)
        merged = self.root / "merged.pdf"
        writer.write(merged)
        writer.close()
        self.assertEqual(len(PdfReader(merged).pages), 2)
        self.assertEqual(len(convert_from_path(source, dpi=72)), 1)
        command("qpdf", "--check", str(merged))

    def test_russian_english_scan_ocr(self):
        from PIL import Image, ImageDraw, ImageFont

        languages = command("tesseract", "--list-langs")
        self.assertTrue({"eng", "rus", "osd"}.issubset(set(languages.split())))
        image = Image.new("RGB", (1600, 600), "white")
        font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 64)
        draw = ImageDraw.Draw(image)
        draw.text((100, 100), "Hermes operator 123", fill="black", font=font)
        draw.text((100, 250), "ПРИВЕТ МИР", fill="black", font=font)
        source = self.root / "scan.pdf"
        image.save(source, resolution=150)
        target = self.root / "searchable.pdf"
        command("ocrmypdf", "--jobs", "1", "--output-type", "pdf", "--optimize", "0",
                "-l", "eng+rus", str(source), str(target), timeout=180)
        text = self.pdf_text(target)
        self.assertIn("Hermes", text)
        self.assertIn("ПРИВЕТ", text)

    def test_data_chart_and_archives(self):
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        import pandas as pd

        frame = pd.DataFrame({"customer": ["Maria", "Maria"], "amount": [12, 30]})
        self.assertEqual(frame.groupby("customer")["amount"].sum()["Maria"], 42)
        frame.plot(y="amount")
        path = self.root / "chart.png"
        plt.savefig(path)
        plt.close("all")
        self.assertGreater(path.stat().st_size, 100)
        archive = self.root / "report.zip"
        command("zip", "-j", str(archive), str(path))
        self.assertIn("chart.png", command("unzip", "-l", str(archive)))
        self.assertIn("Everything is Ok", command("7z", "t", str(archive)))


if __name__ == "__main__":
    Path(os.environ["HOME"]).mkdir(parents=True, exist_ok=True)
    unittest.main(verbosity=2)
